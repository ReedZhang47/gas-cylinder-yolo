"""Draw Fig. 2: provenance and the two generative-data workflows.

The photographs are embedded from the author's original PowerPoint, not linked
to an external file.  All diagrams and type remain vector/editable in the SVG.
The PDF is the manuscript-ready version of the same artwork.
"""

from __future__ import annotations

import base64
import html
import math
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


HERE = Path(__file__).resolve().parent
SOURCE_PPTX = HERE / "fig1_pipeline.pptx"
OUT_PDF = HERE.parent / "fig2_generation_workflows.pdf"
OUT_SVG = HERE / "generation_workflows_v5.svg"
W, H = 760, 466

INK = "#1B3448"
MUTED = "#566C7D"
LINE = "#9EB3C1"
BLUE = "#426F96"
TEAL = "#178A93"
PURPLE = "#755CAD"


def register_fonts() -> tuple[str, str]:
    font_dir = Path(r"C:\Windows\Fonts")
    regular, bold = font_dir / "arial.ttf", font_dir / "arialbd.ttf"
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont("Fig2Arial", str(regular)))
        pdfmetrics.registerFont(TTFont("Fig2ArialBold", str(bold)))
        return "Fig2Arial", "Fig2ArialBold"
    return "Helvetica", "Helvetica-Bold"


FONT, FONT_BOLD = register_fonts()


def source_photo(shape_number: int) -> bytes:
    """Return a compact, print-resolution JPEG from an original PPT picture."""
    shape = Presentation(str(SOURCE_PPTX)).slides[0].shapes[shape_number]
    image = Image.open(BytesIO(shape.image.blob)).convert("RGB")
    image = ImageOps.fit(image, (1000, 750), method=Image.Resampling.LANCZOS)
    out = BytesIO()
    image.save(out, format="JPEG", quality=91, optimize=True)
    return out.getvalue()


class Figure:
    def __init__(self) -> None:
        self.pdf = canvas.Canvas(str(OUT_PDF), pagesize=(W, H), pageCompression=1)
        self.pdf.setTitle("Data provenance and generative workflows")
        self.pdf.setAuthor("Gas Cylinder Safety Detection project")
        self.svg: list[str] = [
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}pt" height="{H}pt" '
            f'viewBox="0 0 {W} {H}" role="img" '
            'aria-label="Common image source, C image-editing workflow, and D LoRA text-to-image workflow">',
            f'<rect x="0" y="0" width="{W}" height="{H}" fill="#ffffff"/>',
        ]

    def box(
        self, x: float, y: float, w: float, h: float, *,
        fill: str = "#ffffff", stroke: str = LINE, radius: float = 7,
        width: float = 1.0, dashed: bool = False,
    ) -> None:
        p = self.pdf
        p.saveState()
        p.setFillColor(HexColor(fill))
        p.setStrokeColor(HexColor(stroke))
        p.setLineWidth(width)
        if dashed:
            p.setDash(4, 3)
        p.roundRect(x, H - y - h, w, h, radius, stroke=1, fill=1)
        p.restoreState()
        dash = ' stroke-dasharray="4 3"' if dashed else ""
        self.svg.append(
            f'<rect x="{x:g}" y="{y:g}" width="{w:g}" height="{h:g}" '
            f'rx="{radius:g}" fill="{fill}" stroke="{stroke}" '
            f'stroke-width="{width:g}"{dash}/>'
        )

    def text(
        self, x: float, y: float, value: str, *, size: float = 11.5,
        color: str = INK, bold: bool = False, align: str = "left",
        max_width: float | None = None,
    ) -> None:
        font = FONT_BOLD if bold else FONT
        measured = pdfmetrics.stringWidth(value, font, size)
        if max_width is not None and measured > max_width + 0.2:
            raise ValueError(f"Text exceeds box ({measured:.1f}>{max_width}): {value}")
        p = self.pdf
        p.saveState()
        p.setFillColor(HexColor(color))
        p.setFont(font, size)
        baseline = H - y - 0.34 * size
        if align == "center":
            p.drawCentredString(x, baseline, value)
        elif align == "right":
            p.drawRightString(x, baseline, value)
        else:
            p.drawString(x, baseline, value)
        p.restoreState()
        anchor = {"left": "start", "center": "middle", "right": "end"}[align]
        weight = "700" if bold else "400"
        self.svg.append(
            f'<text x="{x:g}" y="{y:g}" dominant-baseline="middle" '
            f'text-anchor="{anchor}" font-family="Arial, Helvetica, sans-serif" '
            f'font-size="{size:g}" font-weight="{weight}" fill="{color}">'
            f'{html.escape(value)}</text>'
        )

    def line(
        self, points: list[tuple[float, float]], *, color: str = LINE,
        width: float = 1.15, dashed: bool = False, arrow: bool = False,
    ) -> None:
        p = self.pdf
        p.saveState()
        p.setStrokeColor(HexColor(color))
        p.setLineWidth(width)
        p.setLineJoin(1)
        p.setLineCap(1)
        if dashed:
            p.setDash(4, 3)
        path = p.beginPath()
        path.moveTo(points[0][0], H - points[0][1])
        for x, y in points[1:]:
            path.lineTo(x, H - y)
        p.drawPath(path, stroke=1, fill=0)
        p.restoreState()
        coords = " ".join(f"{x:g},{y:g}" for x, y in points)
        dash = ' stroke-dasharray="4 3"' if dashed else ""
        self.svg.append(
            f'<polyline points="{coords}" fill="none" stroke="{color}" '
            f'stroke-width="{width:g}" stroke-linejoin="round" '
            f'stroke-linecap="round"{dash}/>'
        )
        if arrow:
            x1, y1 = points[-2]
            x2, y2 = points[-1]
            length = math.hypot(x2 - x1, y2 - y1)
            ux, uy = (x2 - x1) / length, (y2 - y1) / length
            px, py = -uy, ux
            base_x, base_y = x2 - 6.2 * ux, y2 - 6.2 * uy
            tri = [
                (x2, y2),
                (base_x + 3.0 * px, base_y + 3.0 * py),
                (base_x - 3.0 * px, base_y - 3.0 * py),
            ]
            p.saveState()
            p.setFillColor(HexColor(color))
            shape = p.beginPath()
            shape.moveTo(tri[0][0], H - tri[0][1])
            for x, y in tri[1:]:
                shape.lineTo(x, H - y)
            shape.close()
            p.drawPath(shape, stroke=0, fill=1)
            p.restoreState()
            vertices = " ".join(f"{x:g},{y:g}" for x, y in tri)
            self.svg.append(f'<polygon points="{vertices}" fill="{color}"/>')

    def photo(self, x: float, y: float, w: float, h: float, jpeg: bytes) -> None:
        self.pdf.drawImage(ImageReader(BytesIO(jpeg)), x, H - y - h,
                           width=w, height=h, mask="auto")
        uri = base64.b64encode(jpeg).decode("ascii")
        self.svg.append(
            f'<image x="{x:g}" y="{y:g}" width="{w:g}" height="{h:g}" '
            f'href="data:image/jpeg;base64,{uri}"/>'
        )
        self.pdf.saveState()
        self.pdf.setStrokeColor(HexColor("#9CB1C1"))
        self.pdf.setLineWidth(0.8)
        self.pdf.rect(x, H - y - h, w, h, stroke=1, fill=0)
        self.pdf.restoreState()
        self.svg.append(
            f'<rect x="{x:g}" y="{y:g}" width="{w:g}" height="{h:g}" '
            'fill="none" stroke="#9CB1C1" stroke-width="0.8"/>'
        )

    def save(self) -> None:
        self.pdf.showPage()
        self.pdf.save()
        self.svg.append("</svg>")
        OUT_SVG.write_text("\n".join(self.svg) + "\n", encoding="utf-8")


def make_figure() -> None:
    f = Figure()
    # Three equally legible, consistently aligned columns.
    f.box(18, 18, 210, 430, fill="#F8FBFD", stroke="#D8E4EB", radius=9)
    f.box(240, 18, 248, 430, fill="#F7FBFB", stroke="#D8E9E9", radius=9)
    f.box(500, 18, 242, 430, fill="#FAF8FD", stroke="#E5DFF0", radius=9)
    f.box(18, 18, 210, 51, fill="#EAF2F8", stroke="#D8E4EB", radius=9)
    f.box(240, 18, 248, 51, fill="#E8F5F5", stroke="#D8E9E9", radius=9)
    f.box(500, 18, 242, 51, fill="#F0EBF7", stroke="#E5DFF0", radius=9)

    # (a) The photographs are true source examples taken from the old PPT.
    f.text(31, 38, "(a)  DATA ORIGIN", size=15.0, bold=True, color=BLUE, max_width=184)
    f.text(31, 57, "shared by both generative arms", size=10.6,
           color=MUTED, max_width=184)
    f.text(32, 90, "real93", size=17.0, bold=True, max_width=177)
    f.text(32, 110, "93 annotated photographs", size=11.8, max_width=177)
    f.photo(32, 132, 86, 65, source_photo(7))
    f.photo(128, 132, 86, 65, source_photo(9))
    f.text(32, 212, "Examples from on-site inspection", size=10.5,
           color=MUTED, max_width=185)
    f.text(32, 234, "Corpus also includes web-sourced", size=10.8,
           max_width=185)
    f.text(32, 249, "real photographs.", size=10.8, max_width=185)
    f.line([(123, 265), (123, 281)], color=BLUE, arrow=True)
    f.box(32, 286, 182, 50, fill="#FFFFFF", stroke=TEAL, radius=6)
    f.text(44, 302, "C  Photo reference", size=12.6, bold=True,
           color=TEAL, max_width=158)
    f.text(44, 322, "edit an existing scene", size=10.8, max_width=158)
    f.box(32, 348, 182, 50, fill="#FFFFFF", stroke=PURPLE, radius=6)
    f.text(44, 364, "D  LoRA training source", size=12.3, bold=True,
           color=PURPLE, max_width=158)
    f.text(44, 384, "learn from the same 93", size=10.8, max_width=158)
    f.text(123, 425, "Shared origin; different synthesis", size=10.5,
           color=MUTED, align="center", max_width=187)

    # (b) Image editing retains the photographed context.
    f.text(253, 38, "(b)  C / IMAGE EDITING", size=14.8, bold=True,
           color=TEAL, max_width=222)
    f.text(253, 57, "complete: 1,085 reviewed images", size=10.6,
           color=MUTED, max_width=222)
    for a, b in ((137, 159), (243, 263), (320, 341)):
        f.line([(364, a), (364, b)], color=TEAL, arrow=True)
    f.box(254, 91, 220, 46, fill="#FFFFFF", stroke="#78ADB0", radius=6)
    f.text(266, 106, "Photograph + edit instruction", size=12.6,
           bold=True, max_width=194)
    f.text(266, 125, "reference scene and target state", size=10.6,
           color=MUTED, max_width=194)
    f.box(254, 159, 220, 84, fill="#FFFFFF", stroke=TEAL, radius=6,
          width=1.15)
    f.text(266, 177, "Qwen-Image-Edit-2509", size=13.2,
           bold=True, color=TEAL, max_width=194)
    f.text(266, 198, "Qwen 2.5 VL encoder + VAE", size=11.0,
           max_width=194)
    f.text(266, 218, "4-step Lightning LoRA mode", size=11.0,
           max_width=194)
    f.box(254, 263, 220, 57, fill="#FFFFFF", stroke="#78ADB0", radius=6)
    f.text(266, 281, "Sample and decode", size=12.6, bold=True,
           max_width=194)
    f.text(266, 302, "edited scene candidate", size=10.8,
           color=MUTED, max_width=194)
    f.box(254, 341, 220, 63, fill="#FFFFFF", stroke=TEAL, radius=6,
          width=1.15)
    f.text(266, 358, "Review image and boxes", size=12.6,
           bold=True, max_width=194)
    f.text(266, 378, "1,085 retained for YOLO", size=11.0,
           max_width=194)
    f.text(364, 426, "Same scene; modified safety state", size=10.5,
           color=MUTED, align="center", max_width=225)

    # (c) LoRA training is done; the final D dataset remains prospective.
    f.text(513, 38, "(c)  D / LORA TEXT-TO-IMAGE", size=13.7,
           bold=True, color=PURPLE, max_width=216)
    f.text(513, 57, "LoRA trained; dataset pending", size=10.6,
           color=MUTED, max_width=216)
    f.line([(621, 148), (621, 165)], color=PURPLE, arrow=True)
    f.line([(621, 239), (621, 257)], color=PURPLE, arrow=True)
    f.line([(621, 320), (621, 340)], color=PURPLE,
           dashed=True, arrow=True)
    f.box(514, 91, 214, 57, fill="#FFFFFF", stroke="#B09CCB", radius=6)
    f.text(526, 108, "real93 + training tags", size=12.8,
           bold=True, max_width=189)
    f.text(526, 129, "Qwen tagging; threshold 0.30", size=10.8,
           max_width=189)
    f.box(514, 165, 214, 74, fill="#FFFFFF", stroke=PURPLE,
          radius=6, width=1.15)
    f.text(526, 182, "Train Qwen-Image LoRA", size=12.8,
           bold=True, color=PURPLE, max_width=189)
    f.text(526, 202, "2,000 steps | LR 4e-4 | rank 16", size=10.7,
           max_width=189)
    f.text(526, 221, "standard workflow; no crop", size=10.7,
           max_width=189)
    f.box(514, 257, 214, 63, fill="#FFFFFF", stroke="#B09CCB", radius=6)
    f.text(526, 275, "Qwen-Image-2.1 + LoRA", size=12.8,
           bold=True, max_width=189)
    f.text(526, 297, "text prompt to candidate scene", size=10.8,
           color=MUTED, max_width=189)
    f.box(514, 340, 214, 64, fill="#FFFFFF", stroke=PURPLE,
          radius=6, width=1.15, dashed=True)
    f.text(526, 357, "Screen and annotate", size=12.8,
           bold=True, max_width=189)
    f.text(526, 380, "target: 1,085 accepted images", size=10.8,
           max_width=189)
    f.text(621, 426, "New-scene coverage to be tested", size=10.5,
           color=MUTED, align="center", max_width=216)

    f.save()
    print(OUT_PDF)
    print(OUT_SVG)


if __name__ == "__main__":
    make_figure()
